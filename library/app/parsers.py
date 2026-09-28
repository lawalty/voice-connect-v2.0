from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document as DocxDocument
from charset_normalizer import from_bytes
from lxml import html
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt", ".html", ".htm"}
SUPPORTED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/markdown",
    "text/plain",
    "text/html",
    "text/x-markdown",
    "application/xhtml+xml",
    "application/octet-stream",
}


class DocumentParseError(ValueError):
    pass


def validate_document(filename: str, mime_type: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            "Supported document types are PDF, DOCX, Markdown, TXT, and HTML"
        )
    if mime_type and mime_type not in SUPPORTED_MIME_TYPES:
        raise DocumentParseError(f"Unsupported media type: {mime_type}")
    return extension


def extract_text(
    filename: str,
    mime_type: str,
    content: bytes,
    *,
    pdf_ocr_enabled: bool = False,
    pdf_ocr_languages: str = "eng",
    pdf_ocr_dpi: int = 200,
    pdf_ocr_min_embedded_chars: int = 32,
    pdf_ocr_page_timeout_seconds: float = 90.0,
) -> str:
    extension = validate_document(filename, mime_type)
    try:
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
            if pdf_ocr_enabled:
                text = _extract_pdf_text_with_ocr(
                    content,
                    page_texts,
                    languages=pdf_ocr_languages,
                    dpi=pdf_ocr_dpi,
                    min_embedded_chars=pdf_ocr_min_embedded_chars,
                    page_timeout_seconds=pdf_ocr_page_timeout_seconds,
                )
            else:
                text = "\n\n".join(page_texts)
        elif extension == ".docx":
            document = DocxDocument(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(paragraphs)
        elif extension in {".html", ".htm"}:
            text = _extract_html(content)
        else:
            text = _decode_text(content)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"Could not parse {filename}: {exc}") from exc
    if not text.strip():
        if extension == ".pdf" and pdf_ocr_enabled:
            raise DocumentParseError(
                "The PDF contains no readable text after local OCR"
            )
        raise DocumentParseError(
            "The document contains no extractable text; scanned PDFs need OCR before ingestion"
        )
    return text


def _extract_pdf_text_with_ocr(
    content: bytes,
    page_texts: list[str],
    *,
    languages: str,
    dpi: int,
    min_embedded_chars: int,
    page_timeout_seconds: float,
) -> str:
    pages_needing_ocr = [
        index
        for index, text in enumerate(page_texts, start=1)
        if _meaningful_character_count(text) < min_embedded_chars
    ]
    if not pages_needing_ocr:
        return "\n\n".join(page_texts)

    if not re.fullmatch(r"[A-Za-z0-9_.+-]+", languages):
        raise DocumentParseError("RAG_OCR_LANGUAGES contains invalid characters")
    if not 100 <= dpi <= 400:
        raise DocumentParseError("RAG_OCR_DPI must be between 100 and 400")
    if page_timeout_seconds <= 0:
        raise DocumentParseError("RAG_OCR_PAGE_TIMEOUT_SECONDS must be positive")

    pdftoppm = shutil.which("pdftoppm")
    tesseract = shutil.which("tesseract")
    if not pdftoppm or not tesseract:
        raise DocumentParseError(
            "Local PDF OCR is unavailable because pdftoppm or Tesseract is missing"
        )

    merged = list(page_texts)
    with tempfile.TemporaryDirectory(prefix="hermes-rag-ocr-") as temp_dir:
        temp_root = Path(temp_dir)
        pdf_path = temp_root / "source.pdf"
        pdf_path.write_bytes(content)
        for page_number in pages_needing_ocr:
            output_prefix = temp_root / f"page-{page_number}"
            image_path = output_prefix.with_suffix(".png")
            _run_ocr_command(
                [
                    pdftoppm,
                    "-f",
                    str(page_number),
                    "-l",
                    str(page_number),
                    "-r",
                    str(dpi),
                    "-png",
                    "-singlefile",
                    str(pdf_path),
                    str(output_prefix),
                ],
                operation=f"render PDF page {page_number}",
                timeout_seconds=page_timeout_seconds,
            )
            completed = _run_ocr_command(
                [
                    tesseract,
                    str(image_path),
                    "stdout",
                    "-l",
                    languages,
                    "--psm",
                    "3",
                ],
                operation=f"OCR PDF page {page_number}",
                timeout_seconds=page_timeout_seconds,
            )
            ocr_text = completed.stdout.strip()
            if ocr_text:
                merged[page_number - 1] = ocr_text
            image_path.unlink(missing_ok=True)
    return "\n\n".join(text for text in merged if text.strip())


def _meaningful_character_count(value: str) -> int:
    return sum(character.isalnum() for character in value)


def _run_ocr_command(
    command: list[str], *, operation: str, timeout_seconds: float
) -> subprocess.CompletedProcess[str]:
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            check=False,
            encoding="utf-8",
            errors="replace",
            timeout=timeout_seconds,
        )
    except subprocess.TimeoutExpired as exc:
        raise DocumentParseError(
            f"Local OCR timed out while trying to {operation}"
        ) from exc
    if completed.returncode != 0:
        detail = completed.stderr.strip()[:600]
        raise DocumentParseError(
            f"Local OCR could not {operation}: {detail or 'command failed'}"
        )
    return completed


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        match = from_bytes(content).best()
        if match is None or match.percent_chaos > 20:
            raise DocumentParseError(
                "Text encoding could not be detected safely"
            )
        return str(match)


def _extract_html(content: bytes) -> str:
    document = html.fromstring(content)
    for node in document.xpath(
        "//script|//style|//noscript|//svg|//template|//nav|//footer|//aside|//form"
    ):
        node.drop_tree()
    roots = document.xpath("//main|//article|//*[@role='main']")
    root = roots[0] if roots else document
    blocks: list[str] = []
    title_nodes = document.xpath("//title")
    if title_nodes:
        title = _clean_html_text(title_nodes[0].text_content())
        if title:
            blocks.append(title)
    for node in root.xpath(".//h1|.//h2|.//h3|.//h4|.//h5|.//h6|.//p|.//li|.//dt|.//dd|.//blockquote|.//pre|.//th|.//td"):
        value = _clean_html_text(node.text_content())
        if value and (not blocks or value != blocks[-1]):
            if node.tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                value = f"{'#' * int(node.tag[1])} {value}"
            blocks.append(value)
    if not blocks:
        fallback = _clean_html_text(root.text_content())
        if fallback:
            blocks.append(fallback)
    return "\n\n".join(blocks)


def _clean_html_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()

